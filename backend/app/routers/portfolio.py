from __future__ import annotations

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy.orm import Session
from concurrent.futures import ThreadPoolExecutor, as_completed
import csv
import io
import openpyxl
from app.models.database import get_db, PortfolioPosition
from app.services.stock_analyzer import get_stock_analysis
from app.dependencies.auth import get_current_user

router = APIRouter()

# Symbols to skip entirely (unresolvable Fidelity placeholders)
SKIP_SYMBOLS = {"FCASH**", "FDRXX**", "SPAXX**", "FDRXX**"}

# Cash / money-market symbols — import with value but skip yfinance analysis
CASH_SYMBOLS = {"SPAXX", "FDRXX", "FCASH", "CORE**", "MMDA1", "MMDA4", "SWEEP"}


def _parse_money(val):
    """Strip broker formatting ($, +, %, commas, parens for negatives) and return float or None."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    if not val or val.strip() in ("--", "", "N/A", "n/a"):
        return None
    cleaned = val.strip().replace("$", "").replace(",", "").replace("+", "").replace("%", "")
    # Handle parenthetical negatives like (123.45)
    if cleaned.startswith("(") and cleaned.endswith(")"):
        cleaned = "-" + cleaned[1:-1]
    try:
        return float(cleaned)
    except ValueError:
        return None


def _detect_broker(headers: list[str]) -> str:
    """Detect broker from CSV column headers."""
    header_set = {h.strip() for h in headers}
    if "Market Value" in header_set or "Unit Cost" in header_set:
        return "etrade"
    if "Purchased Qty." in header_set or "Est. Market Value" in header_set:
        return "etrade_plan"
    return "fidelity"


def _parse_etrade_plan_row(row: dict) -> dict | None:
    """Parse one row from an E*Trade stock plan ByBenefitType CSV."""
    if row.get("Record Type", "").strip() != "Purchase":
        return None
    symbol = row.get("Symbol", "").strip()
    if not symbol:
        return None

    net_shares = _parse_money(row.get("Net Shares", "") or row.get("Purchased Qty.", ""))
    if net_shares is None or net_shares <= 0:
        return None

    current_value = _parse_money(row.get("Est. Market Value", ""))
    avg_cost = _parse_money(row.get("Est. Cost Basis (per share):", ""))
    cost_basis_total = round(net_shares * avg_cost, 2) if net_shares and avg_cost else None
    total_gl_dollar = _parse_money(row.get("Expected Gain/Loss", ""))
    total_gl_pct = None
    if cost_basis_total and cost_basis_total > 0 and total_gl_dollar is not None:
        total_gl_pct = round(total_gl_dollar / cost_basis_total * 100, 2)
    last_price = round(current_value / net_shares, 4) if current_value and net_shares else None
    purchase_date = row.get("Purchase Date", "").strip()

    return dict(
        broker="etrade",
        account_number="etrade_plan",
        account_name=f"E*Trade Plan ({purchase_date})" if purchase_date else "E*Trade Plan",
        ticker=symbol.upper(),
        company_name="",
        shares=net_shares,
        avg_cost=avg_cost,
        cost_basis_total=cost_basis_total,
        last_price=last_price,
        current_value=current_value,
        total_gl_dollar=total_gl_dollar,
        total_gl_pct=total_gl_pct,
    )


def _parse_fidelity_row(row: dict) -> dict | None:
    """Parse one row from a Fidelity positions CSV."""
    symbol = row.get("Symbol", "").strip()
    if not symbol or symbol in SKIP_SYMBOLS:
        return None
    if not row.get("Account Number", "").strip():
        return None

    if "Pending" in symbol:
        current_value = _parse_money(row.get("Current Value", ""))
        if current_value is None:
            return None
        return dict(
            broker="fidelity",
            account_number=row.get("Account Number", "").strip(),
            account_name=row.get("Account Name", "").strip(),
            ticker="PENDING",
            company_name="Pending Activity",
            shares=0,
            avg_cost=None,
            cost_basis_total=None,
            last_price=None,
            current_value=current_value,
            total_gl_dollar=None,
            total_gl_pct=None,
        )

    is_cash = symbol in CASH_SYMBOLS
    shares = _parse_money(row.get("Quantity", ""))
    current_value = _parse_money(row.get("Current Value", ""))

    if is_cash:
        if not current_value:
            return None
    else:
        if shares is None or shares <= 0:
            return None

    return dict(
        broker="fidelity",
        account_number=row.get("Account Number", "").strip(),
        account_name=row.get("Account Name", "").strip(),
        ticker=symbol.upper(),
        company_name=row.get("Description", "").strip() or ("Cash & Money Market" if is_cash else ""),
        shares=shares,
        avg_cost=_parse_money(row.get("Average Cost Basis", "")),
        cost_basis_total=_parse_money(row.get("Cost Basis Total", "")),
        last_price=_parse_money(row.get("Last Price", "")),
        current_value=current_value,
        total_gl_dollar=_parse_money(row.get("Total Gain/Loss Dollar", "")),
        total_gl_pct=_parse_money(row.get("Total Gain/Loss Percent", "")),
    )


def _parse_etrade_row(row: dict, account_label: str) -> dict | None:
    """Parse one row from an E*Trade positions CSV."""
    symbol = row.get("Symbol", "").strip()
    if not symbol or symbol in SKIP_SYMBOLS or "Pending" in symbol:
        return None

    is_cash = symbol in CASH_SYMBOLS
    shares = _parse_money(row.get("Quantity", ""))
    current_value = _parse_money(row.get("Market Value", ""))

    if is_cash:
        if not current_value:
            return None
    else:
        if shares is None or shares <= 0:
            return None

    avg_cost = _parse_money(row.get("Unit Cost", ""))
    total_gl_dollar = _parse_money(row.get("Total Gain ($)", "") or row.get("Total Gain", ""))
    total_gl_pct = _parse_money(row.get("Total Gain (%)", "") or row.get("Total Gain %", ""))

    # Derive cost basis total if not directly available
    cost_basis_total = None
    if current_value is not None and total_gl_dollar is not None:
        cost_basis_total = current_value - total_gl_dollar
    elif shares and avg_cost:
        cost_basis_total = shares * avg_cost

    return dict(
        broker="etrade",
        account_number=account_label,
        account_name=account_label or "E*Trade",
        ticker=symbol.upper(),
        company_name=row.get("Description", "").strip() or ("Cash & Money Market" if is_cash else ""),
        shares=shares,
        avg_cost=avg_cost,
        cost_basis_total=cost_basis_total,
        last_price=_parse_money(row.get("Last Price", "")),
        current_value=current_value,
        total_gl_dollar=total_gl_dollar,
        total_gl_pct=total_gl_pct,
    )


def _xlsx_row_to_dict(headers: list, raw_row: tuple) -> dict:
    """Convert an openpyxl row tuple to a string-keyed dict, keeping first non-None for duplicate headers."""
    result: dict = {}
    for i, v in enumerate(raw_row):
        if i >= len(headers):
            break
        key = headers[i]
        if key not in result or result[key] is None:
            result[key] = str(v) if isinstance(v, str) else v
    return result


def _parse_rsu_row(row: dict) -> dict | None:
    """Parse a Grant row from the Restricted Stock sheet of an E*Trade XLSX."""
    if str(row.get("Record Type", "")).strip() != "Grant":
        return None
    symbol = str(row.get("Symbol", "") or "").strip()
    if not symbol:
        return None
    shares = _parse_money(row.get("Granted Qty."))
    current_value = _parse_money(row.get("Est. Market Value"))
    if not shares or not current_value:
        return None
    grant_date = str(row.get("Grant Date", "") or "").strip()
    return dict(
        broker="etrade",
        account_number="etrade_plan",
        account_name=f"E*Trade Plan (RSU {grant_date})" if grant_date else "E*Trade Plan (RSU)",
        ticker=symbol.upper(),
        company_name="",
        shares=shares,
        avg_cost=0.0,
        cost_basis_total=0.0,
        last_price=round(current_value / shares, 4) if shares else None,
        current_value=current_value,
        total_gl_dollar=current_value,
        total_gl_pct=None,
    )


def _parse_etrade_plan_xlsx(content_bytes: bytes) -> list[dict]:
    """Parse all sheets from an E*Trade stock plan XLSX and return position dicts."""
    wb = openpyxl.load_workbook(io.BytesIO(content_bytes), read_only=True, data_only=True)
    results = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        sheet_name = ws.title.lower()
        for raw_row in rows[1:]:
            row = _xlsx_row_to_dict(headers, raw_row)
            if "restricted" in sheet_name:
                parsed = _parse_rsu_row(row)
            else:
                parsed = _parse_etrade_plan_row(row)
            if parsed:
                results.append(parsed)
    return results


@router.post("/detect")
async def detect_broker(
    file: UploadFile = File(...),
    user_id: str = Depends(get_current_user),
):
    """Sniff the first line of the CSV and return detected broker."""
    first_chunk = await file.read(2048)
    first_line = first_chunk.decode("utf-8-sig").splitlines()[0]
    headers = [h.strip() for h in first_line.split(",")]
    broker = _detect_broker(headers)
    return {"broker": broker}


@router.post("/import")
async def import_portfolio(
    file: UploadFile = File(...),
    account_label: str = Form(""),
    db: Session = Depends(get_db),
    user_id: str = Depends(get_current_user),
):
    MAX_FILE_BYTES = 5 * 1024 * 1024  # 5 MB
    content_bytes = await file.read(MAX_FILE_BYTES + 1)
    if len(content_bytes) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="File too large — 5 MB limit")

    filename = (file.filename or "").lower()
    if filename.endswith(".xlsx"):
        positions_data = _parse_etrade_plan_xlsx(content_bytes)
        broker = "etrade_plan"
    else:
        content = content_bytes.decode("utf-8-sig")
        clean_lines = [l for l in content.splitlines() if not l.startswith('"')]
        reader = csv.DictReader(io.StringIO("\n".join(clean_lines)))
        headers = reader.fieldnames or []
        broker = _detect_broker(list(headers))
        positions_data = []
        for row in reader:
            if broker == "etrade":
                parsed = _parse_etrade_row(row, account_label)
            elif broker == "etrade_plan":
                parsed = _parse_etrade_plan_row(row)
            else:
                parsed = _parse_fidelity_row(row)
            if parsed:
                positions_data.append(parsed)

    # Smart clear: replace only the imported broker's data for this user
    if broker == "etrade" and account_label:
        db.query(PortfolioPosition).filter(
            PortfolioPosition.user_id == user_id,
            PortfolioPosition.broker == "etrade",
            PortfolioPosition.account_name == (account_label or "E*Trade"),
        ).delete()
    elif broker == "etrade_plan":
        db.query(PortfolioPosition).filter(
            PortfolioPosition.user_id == user_id,
            PortfolioPosition.broker == "etrade",
            PortfolioPosition.account_number == "etrade_plan",
        ).delete()
    else:
        db.query(PortfolioPosition).filter(
            PortfolioPosition.user_id == user_id,
            PortfolioPosition.broker == broker,
        ).delete()

    db.add_all([PortfolioPosition(user_id=user_id, **p) for p in positions_data])
    db.commit()

    return {
        "imported": len(positions_data),
        "broker": broker,
        "message": f"Imported {len(positions_data)} positions from {broker.title()}",
    }


@router.get("/")
def get_portfolio(db: Session = Depends(get_db), user_id: str = Depends(get_current_user)):
    positions = db.query(PortfolioPosition).filter(PortfolioPosition.user_id == user_id).order_by(PortfolioPosition.current_value.desc()).all()

    unique_tickers = list({p.ticker for p in positions if p.ticker not in CASH_SYMBOLS})
    analyses: dict[str, dict] = {}
    with ThreadPoolExecutor(max_workers=10) as pool:
        futures = {pool.submit(get_stock_analysis, ticker): ticker for ticker in unique_tickers}
        for future in as_completed(futures):
            ticker = futures[future]
            try:
                result = future.result()
                analyses[ticker] = result if "error" not in result else None
            except Exception:
                analyses[ticker] = None

    return [
        {
            "id": p.id,
            "broker": p.broker or "fidelity",
            "account_number": p.account_number,
            "account_name": p.account_name,
            "ticker": p.ticker,
            "company_name": p.company_name,
            "shares": p.shares,
            "avg_cost": p.avg_cost,
            "cost_basis_total": p.cost_basis_total,
            "last_price": p.last_price,
            "current_value": p.current_value,
            "total_gl_dollar": p.total_gl_dollar,
            "total_gl_pct": p.total_gl_pct,
            "imported_at": p.imported_at.isoformat() if p.imported_at else None,
            "analysis": analyses.get(p.ticker),
        }
        for p in positions
    ]


@router.get("/summary")
def get_summary(db: Session = Depends(get_db), user_id: str = Depends(get_current_user)):
    positions = db.query(PortfolioPosition).filter(PortfolioPosition.user_id == user_id).all()
    total_value = sum(p.current_value or 0 for p in positions)
    total_gl = sum(p.total_gl_dollar or 0 for p in positions)
    total_cost = sum(p.cost_basis_total or 0 for p in positions)
    total_gl_pct = (total_gl / total_cost * 100) if total_cost > 0 else 0

    accounts: dict = {}
    for p in positions:
        key = p.account_name
        if key not in accounts:
            accounts[key] = {"value": 0, "gl": 0, "count": 0, "broker": p.broker or "fidelity"}
        accounts[key]["value"] += p.current_value or 0
        accounts[key]["gl"] += p.total_gl_dollar or 0
        accounts[key]["count"] += 1

    return {
        "total_value": round(total_value, 2),
        "total_gl_dollar": round(total_gl, 2),
        "total_gl_pct": round(total_gl_pct, 2),
        "total_cost": round(total_cost, 2),
        "position_count": len(positions),
        "accounts": {k: {kk: round(vv, 2) if isinstance(vv, float) else vv for kk, vv in v.items()} for k, v in accounts.items()},
    }


@router.delete("/")
def clear_portfolio(db: Session = Depends(get_db), user_id: str = Depends(get_current_user)):
    db.query(PortfolioPosition).filter(PortfolioPosition.user_id == user_id).delete()
    db.commit()
    return {"message": "Portfolio cleared"}
